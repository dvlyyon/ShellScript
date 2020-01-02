#!/usr/bin/python3

import sys
import re
from enum import IntEnum
from openpyxl import Workbook
import datetime

from openpyxl.chart import (LineChart, Reference,)
from openpyxl.chart.axis import DateAxis

class State(IntEnum):
    INIT = 1
    DATE = 2
    MF_RSS = 3
    MF_DATA = 4
    MF_THREAD = 5
    SMF_RSS = 6
    SMF_DATA = 7
    SMF_THREAD = 8
    MF_TOP = 9
    SMF_TOP = 10
    MF_HEAP = 11

origFile = open(sys.argv[1])

wb = Workbook()
ws = wb.active
ws.title = "Memory Statistics"

title = ["Time", "MF VmData", "SMF VmData", "MF Memory %",  "SMF Memory %", "MF CPU %", "SMF CPU %", "MF VmRSS", "MF Heap", "MF Threads"]
data=[]
data.append(title)
state=State.INIT
tmpline=[0,0,0,0,0,0,0,0,0,0]
topmark=''
heap=0
for line in origFile:
    if state == State.INIT:
        match = re.match("^date$", line)
        if match:
            state=State.DATE
    elif state == State.DATE:
        timefields = re.split('[ ]+',line.strip())
        datestr = timefields[1]+' '+timefields[2]+' '+ \
                  timefields[3]+' '+timefields[5]
        print(datestr)
        date = datetime.datetime.strptime(datestr,'%b %d %H:%M:%S %Y')
        tmpline[0]=date
        state = State.MF_RSS
    elif state == State.MF_RSS:
        match = re.match('^VmRSS:\s+(\d+)\s+',line)
        if match:
            print('Match MF_VmRss')
            tmpline[7]=int(match.group(1))
            state=State.MF_DATA
    elif state == State.MF_DATA:
        match = re.match('^VmData:\s+(\d+)\s+',line)
        if match:
            print('Match MF_DATA')
            tmpline[1]=int(match.group(1))
            state=State.MF_THREAD
    elif state == State.MF_THREAD:
        match = re.match('^Threads:\s+(\d+)',line)
        if match:
            print('Match MF_Thread')
            tmpline[9]=int(match.group(1))
            state=State.SMF_DATA
    elif state == State.SMF_DATA:
        match = re.match('^VmData:\s+(\d+)\s+',line)
        if match:
            print('Match SMF_DATA')
            tmpline[2]=int(match.group(1))
            state=State.MF_TOP
    elif state == State.MF_TOP:
        match1 = re.match('^.*\s+([0-9.]+)+\s+([0-9.]+)\s+[0-9:.]+\s+MF.bin',line)
        match2 = re.match('^.*\s+([0-9.]+)+\s+([0-9.]+)\s+[0-9:.]+\s+SMF.bin',line)
        if match1 or match2:
            print('Match Top')
            if match1:
                topmark=topmark+"X"
                tmpline[3]=float(match1.group(2))
                tmpline[5]=float(match1.group(1))
            elif match2:
                topmark=topmark+"S"
                tmpline[4]=float(match2.group(2))
                tmpline[6]=float(match2.group(1))
            if topmark.find('X')>=0 and topmark.find('S')>=0:
                state=State.MF_HEAP
            elif topmark.find('X')>=0:
                state=State.SMF_TOP
            elif topmark.find('S')>=0:
                state=State.MF_TOP
    elif state == State.SMF_TOP:
        match1 = re.match('^.*\s+([0-9.]+)+\s+([0-9.]+)\s+[0-9:.]+\s+MF.bin',line)
        match2 = re.match('^.*\s+([0-9.]+)+\s+([0-9.]+)\s+[0-9:.]+\s+SMF.bin',line)
        if match1 or match2:
            if match1:
                topmark=topmark+"X"
                tmpline[3]=float(match1.group(2))
                tmpline[5]=float(match1.group(1))
            elif match2:
                topmark=topmark+"S"
                tmpline[4]=float(match2.group(2))
                tmpline[6]=float(match2.group(1))
            if topmark.find('X')>=0 and topmark.find('S')>=0:
                state=State.MF_HEAP
                topmark=''
            elif topmark.find('X')>=0:
                state=State.SMF_TOP
            elif topmark.find('S')>=0:
                state=State.MF_TOP
    elif state == State.MF_HEAP:
        match1 = re.match('^([0-9A-Fa-f]+)-([0-9A-Fa-f]+)\s.*\[heap\]$',line)
        match2 = re.match("^date$", line)
        if match1:
            heap = heap + int(match1.group(2),base=16)-int(match1.group(1),base=16)
            state = State.MF_HEAP
        elif match2:
            tmpline[8]=heap//1024
            heap=0
            data.append(tmpline)
            print(tmpline)
            state = State.DATE
            
for row in data:
    ws.append(row)

data1 = Reference(ws,min_col=2, min_row=1, max_col=3, max_row=len(data))

c1 = LineChart()
c1.title = "Memory Information"
c1.style = 13
c1.y_axis.majorGridlines = None
c1.y_axis.title = "Memory (KB)"
c1.x_axis.title = "Time"

c1.add_data(data1, titles_from_data=True)

s1 = c1.series[0]
s1.graphicalProperties.line.solidFill = "FF0000"
#s1.graphicalProperties.line.dashStyle = "sysDot"
#s1.graphicalProperties.line.width = 100050 # width in EMUs

s2 = c1.series[1]
s2.graphicalProperties.line.solidFill = "990000"
s2.smooth = True

c2 = LineChart()
data2 = Reference(ws, min_col=4, min_row=1, max_col=7, max_row=len(data))
# c2.grouping = "percentStacked"
c2.add_data(data2, titles_from_data=True)
c2.y_axis.axId =200
c2.x_axis.title = "Time"
c2.y_axis.title = "CPU and Memory Ratio"
s21 = c2.series[0]
s21.graphicalProperties.line.solidFill = "0000FF"
#s21.graphicalProperties.line.dashStyle = "sysDash"
s21.smooth=True
s21.graphicalProperties.line.width = 30000 # width in EMUs
s22 = c2.series[1]
s22.graphicalProperties.line.solidFill = "000099"
#s22.graphicalProperties.line.dashStyle = "sysDash"
s22.graphicalProperties.line.width = 30000 # width in EMUs
s23 = c2.series[2]
s23.graphicalProperties.line.solidFill = "00FF00"
#s23.graphicalProperties.line.dashStyle = "sysDash"
s23.graphicalProperties.line.width = 30000 # width in EMUs
s24 = c2.series[3]
s24.graphicalProperties.line.solidFill = "009900"
#s24.graphicalProperties.line.dashStyle = "sysDash"
s24.graphicalProperties.line.width = 30000 # width in EMUs

c2.y_axis.crosses = "max"
c1 += c2
ws.add_chart(c1, "A"+str(len(data)+5))

data1x = Reference(ws,min_col=8, min_row=1, max_col=9, max_row=len(data))

c1x = LineChart()
c1x.title = "RSS Heap and Thread Information"
c1x.style = 13
c1x.y_axis.majorGridlines = None
c1x.y_axis.title = "Memory (KB)"
c1x.x_axis.title = "Time"

c1x.add_data(data1x, titles_from_data=True)

s1x = c1x.series[0]
s1x.graphicalProperties.line.solidFill = "FF0000"

s2x = c1x.series[1]
s2x.graphicalProperties.line.solidFill = "990000"
s2x.smooth = True

c2x = LineChart()
data2x = Reference(ws, min_col=10, min_row=1, max_col=10, max_row=len(data))
c2x.add_data(data2x, titles_from_data=True)
c2x.y_axis.axId =300
c2x.x_axis.title = "Time"
c2x.y_axis.title = "Thread"
s21x = c2x.series[0]
s21x.graphicalProperties.line.solidFill = "0000FF"
#s21x.graphicalProperties.line.dashStyle = "sysDash"
s21x.graphicalProperties.line.width = 30000 # width in EMUs

c2x.y_axis.crosses = "max"
c1x += c2x
ws.add_chart(c1x, "A"+str(len(data)+25))

wb.save(sys.argv[1]+".telemetry.xlsx")
origFile.close()
